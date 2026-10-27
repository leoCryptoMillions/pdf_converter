"""Local Word fidelity evidence, not a human-adjudicated F0 quality gate.

python -m benchmarks.validate_word prepare --output .generated/word-validation/baseline
powershell -File benchmarks/render_word.ps1 -Batch <output>/batch.json
python -m benchmarks.validate_word compare --output <output>

Only IDs/metrics are printed. PDFs, text, images and Word artifacts stay local.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
import html
from io import BytesIO
import json
import math
from pathlib import Path
import statistics
import time

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import pdfplumber

from benchmarks.common import load_manifest
from packages.converters.digital import extract_digital
from packages.exporters.docx import export_docx

DEFAULT_IDS = ['cep_001', 'carta_privacidad_001', 'doc_texto_002']


def visible_characters(text: str) -> Counter:
    """No case, punctuation or Unicode normalization: only whitespace omitted."""
    return Counter(char for char in text if not char.isspace())


def content_comparison(expected: str, actual: str) -> dict:
    left, right = visible_characters(expected), visible_characters(actual)
    missing, extra = left - right, right - left
    return {'expected_nonspace': sum(left.values()), 'actual_nonspace': sum(right.values()),
            'missing_characters': sum(missing.values()), 'extra_characters': sum(extra.values()),
            'exact_character_multiset': left == right,
            'missing_codepoints': {f'U+{ord(k):04X}': v for k, v in sorted(missing.items())},
            'extra_codepoints': {f'U+{ord(k):04X}': v for k, v in sorted(extra.items())},
            'scope': 'Character multiset only; does not verify reading order or semantic correctness.'}


def word_text(word) -> str:
    # Traverse XML once: merged cells and nested tables must not duplicate text.
    return '\n'.join(Paragraph(p, word._body).text for p in word.element.body.iter(qn('w:p')))


def geometry_comparison(source, rendered) -> dict:
    """Describe positions of uniquely matching words; never call this visual F1."""
    def unique(page):
        words = page.extract_words()
        counts = Counter(word['text'] for word in words)
        return {w['text']: w for w in words if counts[w['text']] == 1}
    left, right = unique(source), unique(rendered)
    distances = []
    for text in left.keys() & right.keys():
        a, b = left[text], right[text]
        distances.append(math.hypot(a['x0'] - b['x0'], a['top'] - b['top']))
    distances.sort()
    return {'uniquely_matched_words': len(distances),
            'median_displacement_pt': round(statistics.median(distances), 3) if distances else None,
            'p95_displacement_pt': round(distances[math.ceil(.95 * len(distances)) - 1], 3) if distances else None,
            'scope': 'Unique matching words only; excludes missing/duplicate words and is not a gate score.'}


def pdfium_text(source) -> str:
    import pypdfium2
    texts = []
    with pypdfium2.PdfDocument(source) as pdf:
        for page in pdf:
            try:
                textpage = page.get_textpage()
                try:
                    texts.append(textpage.get_text_range())
                finally:
                    textpage.close()
            finally:
                page.close()
    return '\n'.join(texts)


def prepare(output: Path, file_ids=None) -> dict:
    output = output.resolve()
    manifest = {row.file_id: row for row in load_manifest()}
    output.mkdir(parents=True, exist_ok=True)
    batch = {'created_at': datetime.now(timezone.utc).isoformat(), 'cases': [],
             'scope': 'Local engineering validation; human-adjudicated references remain pending.'}
    for file_id in file_ids or DEFAULT_IDS:
        row = manifest[file_id]
        if not row.authorized or not row.pdf_path.is_file():
            raise ValueError(f'Unauthorized or unavailable case: {file_id}')
        case = output / file_id
        case.mkdir(exist_ok=False)  # Keep earlier runs intact.
        started = time.monotonic()
        ir = extract_digital(row.pdf_path, preserve_layout=True)
        data = export_docx(ir)
        (case / 'result.docx').write_bytes(data)
        (case / 'document.ir.json').write_text(ir.model_dump_json(indent=2), encoding='utf-8')
        word = Document(BytesIO(data))
        actual = word_text(word)
        independent = pdfium_text(row.pdf_path)
        (case / 'source-pdfium.txt').write_text(independent, encoding='utf-8')
        (case / 'word-text.txt').write_text(actual, encoding='utf-8')
        # Independent references are extraction comparisons, never automatically
        # promoted to human-reviewed ground truth or CER acceptance.
        report = {'id': file_id, 'source_sha256': ir.source_hash,
            'source_pages': ir.page_count, 'docx_sections': len(word.sections),
            'docx_tables': len(word.tables),
            'source_to_docx': content_comparison(ir.get_all_text(), actual),
            'pdfium_to_docx': content_comparison(independent, actual),
            'images_by_page': [len(page.images) for page in ir.pages],
            'warnings': [{'page': w.page_number, 'code': w.code} for page in ir.pages for w in page.warnings],
            'engines': [e.model_dump() for e in ir.engines_used],
            'processing_seconds': round(time.monotonic() - started, 3),
            'human_adjudication': 'pending', 'visual_review': 'pending'}
        # Exercise edit/save/reopen through python-docx as well as real Word.
        target = next((p for p in word.paragraphs if p.text.strip()), None)
        if target is None:
            target = word.add_paragraph()
        target.add_run(' [F0_EDIT_TEST]')
        saved = BytesIO()
        word.save(saved)
        report['python_edit_save_reopen'] = '[F0_EDIT_TEST]' in word_text(Document(BytesIO(saved.getvalue())))
        with pdfplumber.open(row.pdf_path) as pdf:
            for number, page in enumerate(pdf.pages, 1):
                page.to_image(resolution=110).original.save(case / f'source-{number}.png')
        (case / 'prepare.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
        batch['cases'].append({'id': file_id, 'docx': str(case / 'result.docx'),
                              'pdf': str(case / 'word-render.pdf'),
                              'edited_docx': str(case / 'edited-in-word.docx'),
                              'word_report': str(case / 'word-runtime.json')})
        print(json.dumps({'id': file_id, 'source_pages': ir.page_count,
                          'missing_chars': report['source_to_docx']['missing_characters'],
                          'extra_chars': report['source_to_docx']['extra_characters']}))
    (output / 'batch.json').write_text(json.dumps(batch, indent=2), encoding='utf-8')
    return batch


def compare(output: Path) -> dict:
    output = output.resolve()
    batch = json.loads((output / 'batch.json').read_text(encoding='utf-8'))
    manifest = {r.file_id: r for r in load_manifest()}
    notes_path = output / 'review-notes.json'
    notes = json.loads(notes_path.read_text(encoding='utf-8')) if notes_path.exists() else {}
    results = []
    for config in batch['cases']:
        file_id = config['id']
        case = output / file_id
        result = json.loads((case / 'prepare.json').read_text(encoding='utf-8'))
        result['visual_review'] = notes.get(file_id, 'pending')
        source = manifest[file_id].pdf_path
        if sha256(source.read_bytes()).hexdigest() != result['source_sha256']:
            raise ValueError('Source changed after preparation')
        result['word_runtime'] = json.loads((case / 'word-runtime.json').read_text(encoding='utf-8-sig'))
        with pdfplumber.open(source) as original, pdfplumber.open(case / 'word-render.pdf') as rendered:
            result['rendered_pages'] = len(rendered.pages)
            result['page_count_matches'] = len(original.pages) == len(rendered.pages)
            result['rendered_text_vs_docx'] = content_comparison((case / 'word-text.txt').read_text(encoding='utf-8'),
                '\n'.join(p.extract_text() or '' for p in rendered.pages))
            result['pages'] = []
            for number, page in enumerate(rendered.pages, 1):
                page.to_image(resolution=110).original.save(case / f'word-{number}.png')
                item = {'page': number, 'rendered_size_pt': [page.width, page.height]}
                if number <= len(original.pages):
                    source_page = original.pages[number - 1]
                    item['source_size_pt'] = [source_page.width, source_page.height]
                    item['size_matches'] = abs(source_page.width - page.width) < .2 and abs(source_page.height - page.height) < .2
                    item['source_to_render'] = content_comparison(source_page.extract_text() or '', page.extract_text() or '')
                    item['geometry'] = geometry_comparison(source_page, page)
                result['pages'].append(item)
        (case / 'comparison.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
        results.append(result)
        print(json.dumps({'id':file_id,'source_pages':result['source_pages'],
                          'rendered_pages':result['rendered_pages'],
                          'render_missing_chars':result['rendered_text_vs_docx']['missing_characters'],
                          'render_extra_chars':result['rendered_text_vs_docx']['extra_characters']}))
    report = {'created_at':datetime.now(timezone.utc).isoformat(), 'cases':results,
              'human_adjudication':'pending', 'gate_decision':'not_evaluated'}
    root = Path(__file__).resolve().parents[1]
    report['implementation_sha256'] = {name: sha256((root / name).read_bytes()).hexdigest() for name in [
        'packages/converters/digital.py', 'packages/converters/layout.py', 'packages/exporters/docx.py',
        'benchmarks/validate_word.py', 'benchmarks/render_word.ps1']}
    (output / 'report.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    write_review_html(report, output)
    return report


def write_review_html(report: dict, output: Path):
    sections = []
    for case in report['cases']:
        file_id = html.escape(case['id'], quote=True)
        pages = []
        for number in range(1, case['rendered_pages'] + 1):
            source = (f'<img src="{file_id}/source-{number}.png" alt="PDF original, página {number}" loading="lazy">'
                      if number <= case['source_pages'] else '<p>No hay página original correspondiente.</p>')
            pages.append(f'<h3>Página {number}</h3><div class="pair"><figure><figcaption>PDF original</figcaption>{source}</figure>'
                         f'<figure><figcaption>Word renderizado</figcaption><img src="{file_id}/word-{number}.png" '
                         f'alt="Word, página {number}" loading="lazy"></figure></div>')
        sections.append(f'<details open><summary>{file_id} · {case["source_pages"]} páginas de origen / '
                        f'{case["rendered_pages"]} en Word</summary><p><a href="{file_id}/result.docx" download>Descargar Word</a> · '
                        f'<a href="{file_id}/edited-in-word.docx" download>Ejemplo editado y reabierto</a></p>{"".join(pages)}</details>')
    document = '''<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Validación local de Word</title><style>
body{font-family:system-ui,sans-serif;margin:24px;background:#eef2f1;color:#192e2b}main{max-width:1600px;margin:auto}
h1{font-size:28px}details{background:white;border:1px solid #c7d4cf;border-radius:12px;padding:20px;margin:24px 0}
summary{font-weight:700;cursor:pointer;padding:12px}a{color:#146c55}a:focus-visible,summary:focus-visible{outline:3px solid #146c55;outline-offset:4px}
.pair{display:grid;grid-template-columns:1fr 1fr;gap:16px}figure{margin:0}figcaption{padding:12px;font-weight:600}
img{width:100%;height:auto;display:block;border:1px solid #c7d4cf;box-sizing:border-box}@media(max-width:700px){.pair{grid-template-columns:1fr}body{margin:12px}}
</style><main><h1>Comparación PDF / Word</h1><p>Revisión técnica local. Compara texto, imágenes, bordes, colores y posiciones.
Los archivos originales y las vistas permanecen en este equipo. La adjudicación humana de referencias y la decisión del Gate F0 siguen pendientes.</p>
'''+''.join(sections)+'</main></html>'
    (output / 'review.html').write_text(document, encoding='utf-8')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=['prepare','compare'])
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--ids', nargs='+')
    args = parser.parse_args()
    if args.action == 'prepare':
        prepare(args.output, args.ids)
    else:
        compare(args.output)
