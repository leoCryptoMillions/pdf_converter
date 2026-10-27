"""Prepare a private, reviewable pilot using the actual local HTTP workflow.

python -m benchmarks.pilot_local --output .generated/mvp-pilot/run-01
Requires requirements-mvp-dev.txt and authorized PDFs already in the corpus.
Automated checks do not approve document quality or the F0 gate.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from hashlib import sha256
from importlib.metadata import version
from io import BytesIO
import json
from pathlib import Path
import re
import time

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from apps import local_review_ui
from benchmarks.common import load_manifest
from benchmarks.validate_word import content_comparison, pdfium_text, word_text
from packages.document_ir.models import DocumentIR

DEFAULT_IDS = ["cep_001", "carta_privacidad_001", "doc_texto_002",
               "estado_cuenta_006", "estado_cuenta_008"]


def check_exports(ir: DocumentIR, artifacts: dict[str, bytes], source: Path) -> dict:
    """Check transport from IR to Office and compare Word with another extractor."""
    word = Document(BytesIO(artifacts["result.docx"]))
    workbook = load_workbook(BytesIO(artifacts["result.xlsx"]), data_only=False)
    try:
        cells_checked = 0
        cells_match = True
        for index, table in enumerate(ir.get_all_tables(), 1):
            sheet = workbook[f"Table_{index}"]
            for cell in table.cells:
                value = sheet.cell(cell.row_index + 1, cell.col_index + 1).value
                cells_match &= ("" if value is None else value) == cell.content
                cells_checked += 1
        formula_count = sum(cell.data_type == "f" for sheet in workbook
                            for row in sheet for cell in row)
        fallback_matches = None
        if not ir.get_all_tables():
            expected = [line for page in ir.pages for line in page.extracted_text.splitlines()]
            actual = ["" if row[0].value is None else row[0].value for row in workbook["Text"]]
            fallback_matches = actual == expected
        text = word_text(word)
        checks = {
            "word_has_editable_content": bool(text.strip()),
            "word_vs_pdfium": content_comparison(pdfium_text(source), text),
            "xlsx_cells_checked": cells_checked,
            "xlsx_cells_match_ir": cells_match,
            "xlsx_formula_count": formula_count,
            "xlsx_has_provenance": {"_Origin", "_Original", "_Warnings"} <= set(workbook.sheetnames),
            "xlsx_text_fallback_matches_ir": fallback_matches,
            "txt_matches_ir": artifacts["result.txt"].decode("utf-8") == ir.get_all_text(),
            "all_pages_processed": len(ir.pages) == ir.page_count,
        }
        checks["passed"] = all((checks["word_has_editable_content"],
            checks["word_vs_pdfium"]["exact_character_multiset"], cells_match,
            formula_count == 0, checks["xlsx_has_provenance"], checks["txt_matches_ir"],
            checks["all_pages_processed"], fallback_matches is not False))
        return checks
    finally:
        workbook.close()


def prepare(output: Path, file_ids: list[str]) -> dict:
    if len(file_ids) != len(set(file_ids)) or not file_ids:
        raise ValueError("Choose unique document IDs")
    manifest = {row.file_id: row for row in load_manifest()}
    rows = []
    for file_id in file_ids:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", file_id):
            raise ValueError("Invalid corpus ID")
        row = manifest.get(file_id)
        if row is None or not row.authorized or not row.category.startswith("digital_"):
            raise ValueError(f"Unavailable, unauthorized or non-digital case: {file_id}")
        if not row.pdf_path.is_file():
            raise ValueError(f"Missing local PDF: {file_id}")
        rows.append(row)
    output = output.resolve()
    private_root = (Path(__file__).resolve().parents[1] / ".generated").resolve()
    if not output.is_relative_to(private_root) or output == private_root:
        raise ValueError("Pilot output must be a new directory inside .generated/")
    output.mkdir(parents=True, exist_ok=False)
    report = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "scope": "Local single-user digital PDF pilot; OCR disabled; no F0 gate approval.",
        "human_acceptance": "pending",
        "versions": {name: version(name) for name in
                     ("pdfplumber", "pypdfium2", "python-docx", "openpyxl", "fastapi", "httpx")},
        "code_sha256": {name: sha256((private_root.parent / name).read_bytes()).hexdigest() for name in
                        ("apps/local_review_ui.py", "packages/converters/digital.py",
                         "packages/converters/layout.py", "packages/exporters/docx.py",
                         "packages/exporters/xlsx.py", "benchmarks/pilot_local.py")},
        "cases": [],
    }
    previous_root = local_review_ui.OUTPUT_ROOT
    try:
        local_review_ui.OUTPUT_ROOT = output / "jobs"
        with TestClient(local_review_ui.app) as client:
            for row in rows:
                started = time.monotonic()
                case = {"id": row.file_id, "human_acceptance": "pending"}
                report["cases"].append(case)
                try:
                    with row.pdf_path.open("rb") as source:
                        response = client.post("/convert", files={"file": ("source.pdf", source, "application/pdf")},
                            data={"formats": ["docx", "xlsx", "txt"]})
                    if response.status_code != 200:
                        raise ValueError(f"Conversion HTTP {response.status_code}")
                    links = re.findall(r'href="(/download/[a-f0-9]+/[^\"]+)"', response.text)
                    artifacts = {}
                    for link in links:
                        download = client.get(link)
                        if download.status_code != 200:
                            raise ValueError(f"Download HTTP {download.status_code}")
                        artifacts[link.rsplit("/", 1)[1]] = download.content
                    expected = {"result.docx", "result.xlsx", "result.txt", "document.ir.json"}
                    if set(artifacts) != expected:
                        raise ValueError("Incomplete downloads")
                    ir = DocumentIR.model_validate_json(artifacts["document.ir.json"])
                    job_id = links[0].split("/")[2]
                    case_dir = output / row.file_id
                    case_dir.mkdir()
                    for name, data in artifacts.items():
                        (case_dir / name).write_bytes(data)
                    # Preserve the uploaded source already written by the HTTP app.
                    (case_dir / "source.pdf").write_bytes((output / "jobs" / job_id / "source.pdf").read_bytes())
                    source_matches = sha256((case_dir / "source.pdf").read_bytes()).hexdigest() == ir.source_hash
                    checks = check_exports(ir, artifacts, case_dir / "source.pdf")
                    checks["source_hash_matches"] = source_matches
                    checks["passed"] &= source_matches
                    case.update({"source_sha256": ir.source_hash, "pages": ir.page_count,
                        "tables": len(ir.get_all_tables()), "checks": checks,
                        "warning_codes": sorted({w.code for w in ir.warnings +
                            [w for page in ir.pages for w in page.warnings]})})
                except Exception as exc:
                    # Keep other cases reviewable; do not print private document content.
                    case["error_type"] = type(exc).__name__
                case["duration_s"] = round(time.monotonic() - started, 3)
                print(f"{row.file_id}: {'checks_passed' if case.get('checks', {}).get('passed') else 'needs_attention'}; human=pending")
    finally:
        local_review_ui.OUTPUT_ROOT = previous_root
    report["checks_passed"] = all(case.get("checks", {}).get("passed", False) for case in report["cases"])
    (output / "report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    with (output / "revision.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["id", "automated_checks", "text_order", "critical_values", "table_structure",
                         "word_appearance", "edit_save_reopen", "decision", "reviewer", "notes"])
        for case in report["cases"]:
            writer.writerow([case["id"], "pass" if case.get("checks", {}).get("passed") else "review",
                             *(["pending"] * 6), "", ""])
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare the private local Word/Excel pilot")
    parser.add_argument("--output", required=True, type=Path, help="New directory inside .generated/")
    parser.add_argument("--ids", nargs="+", default=DEFAULT_IDS, help="Authorized digital corpus IDs")
    args = parser.parse_args()
    try:
        report = prepare(args.output, args.ids)
    except (ValueError, OSError) as exc:
        parser.exit(1, f"Pilot could not be prepared: {exc}\n")
    print(f"Report: {args.output / 'report.json'}; human acceptance pending")
    if not report["checks_passed"]:
        parser.exit(1, "Some cases need attention; see report.json\n")


if __name__ == "__main__":
    main()
