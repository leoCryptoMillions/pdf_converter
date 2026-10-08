"""Small, reproducible F0 evidence report; deliberately not a Go/No-Go gate.

Run: python -m benchmarks.validate_prototype
Artifacts and reference text stay local. Only metrics are printed.
"""
from __future__ import annotations

import hashlib
from decimal import Decimal, InvalidOperation
from importlib.metadata import version
from io import BytesIO
import json
from pathlib import Path
import time

from benchmarks.common import character_error_rate, load_ground_truth, load_manifest, table_f1
from packages.converters.digital import extract_digital
from packages.exporters.xlsx import export_xlsx

DEFAULT_IDS = ["synthetic_digital_001", "estado_cuenta_008", "estado_cuenta_004",
               "yucash_resolucion_001", "scanned_002", "scanned_006"]


def normalize(text: str) -> str:
    return " ".join(text.split())


def pdfium_text_docx_baseline(source: Path, destination: Path, expected_text: str) -> dict:
    """Independent F0 Word route: PDFium text into editable paragraphs.

    This baseline does not reconstruct tables, images, headings or geometry.
    """
    import pypdfium2
    from docx import Document

    output = Document()
    with pypdfium2.PdfDocument(str(source)) as pdf:
        for index in range(len(pdf)):
            if index:
                output.add_page_break()
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                finally:
                    textpage.close()
            finally:
                page.close()
            for line in text.splitlines():
                output.add_paragraph(line)
    output.core_properties.comments = "F0 text-only baseline: PDFium; no image/signature/layout reproduction."
    output.save(destination)
    reopened = Document(destination)
    actual = normalize("\n".join(p.text for p in reopened.paragraphs))
    reopened.add_paragraph("F0 baseline editability verification")
    buffer = BytesIO()
    reopened.save(buffer)
    return {"route": "pdfium_text_only_docx", "cer_whitespace_normalized": character_error_rate(normalize(expected_text), actual),
            "edit_save_reopen": Document(BytesIO(buffer.getvalue())).paragraphs[-1].text == "F0 baseline editability verification",
            "scope": "text-only; no structure/layout guarantee"}


def ocr_page_text(pdf_path: Path, number: int, *, psm: int = 3, max_pixels: int = 20_000_000,
                  dpi: int = 300, preprocess: str = "none") -> str:
    """Render only an annotated full page at 300 DPI with a bounded pixel budget."""
    import pypdfium2
    import pytesseract
    from benchmarks import ocr_benchmark  # configure native binaries/installed languages

    with pypdfium2.PdfDocument(str(pdf_path)) as pdf:
        page = pdf[number - 1]
        try:
            width, height = page.get_size()
            if dpi <= 0 or dpi > 600 or preprocess not in {"none", "autocontrast", "threshold160"}:
                raise ValueError("Invalid OCR preprocessing options")
            scale = dpi / 72
            if (int(width * scale) + 1) * (int(height * scale) + 1) > max_pixels:
                raise ValueError("Rendered page exceeds OCR pixel budget")
            bitmap = page.render(scale=scale)
            try:
                image = bitmap.to_pil()
                processed = image
                try:
                    if preprocess != "none":
                        from PIL import ImageOps
                        processed = ImageOps.autocontrast(image.convert("L"))
                        if preprocess == "threshold160":
                            gray = processed
                            processed = gray.point(lambda value: 255 if value >= 160 else 0)
                            gray.close()
                    return pytesseract.image_to_string(processed, lang="spa+eng", config=f"--psm {psm}", timeout=120)
                finally:
                    if processed is not image:
                        processed.close()
                    image.close()
            finally:
                bitmap.close()
        finally:
            page.close()


def run(file_ids: list[str] | None = None, output: Path = Path(".generated/f0")) -> dict:
    from openpyxl import load_workbook

    manifest = {row.file_id: row for row in load_manifest()}
    results = []
    output.mkdir(parents=True, exist_ok=True)
    for file_id in file_ids or DEFAULT_IDS:
        row = manifest[file_id]
        if not row.authorized or not row.pdf_path.is_file():
            results.append({"file_id": file_id, "skipped": "unauthorized or missing"})
            continue
        reference = load_ground_truth(file_id)
        if not reference:
            results.append({"file_id": file_id, "skipped": "no verified reference"})
            continue
        digest = hashlib.sha256(row.pdf_path.read_bytes()).hexdigest()
        if reference.get("source_sha256", digest) != digest:
            raise ValueError(f"Reference hash mismatch: {file_id}")
        record = {"file_id": file_id, "category": row.category,
                  "reference_scope": reference.get("scope", "legacy_tables"), "source_sha256": digest,
                  "reference_revision": reference.get("reference_revision", 1)}
        selected_pages = [p["page_number"] for p in reference["pages"]] if reference.get("scope", "").startswith("selected_pages_") else None
        record["annotated_pages"] = [p["page_number"] for p in reference["pages"]]
        directory = output / file_id
        directory.mkdir(exist_ok=True)
        document = extract_digital(row.pdf_path, page_numbers=selected_pages)
        record["source_page_count"] = document.page_count
        (directory / "document.ir.json").write_text(document.model_dump_json(indent=2), encoding="utf-8")
        expected_tables = [t for p in reference["pages"] for t in p.get("tables", [])]
        actual_tables = [[[next((c.content for c in table.cells if c.row_index == r and c.col_index == col), "")
                            for col in range(table.columns)] for r in range(table.rows)]
                         for table in document.get_all_tables()]
        record["tables_detected"] = len(actual_tables)
        if expected_tables:
            record["pdfplumber_row_f1"] = table_f1(expected_tables, actual_tables)
            started = time.monotonic()
            try:
                import camelot
                pages_option = ",".join(map(str, sorted(selected_pages))) if selected_pages else "all"
                extracted = [table.df.values.tolist() for table in camelot.read_pdf(str(row.pdf_path), pages=pages_option)]
                record["camelot_row_f1"] = table_f1(expected_tables, extracted)
                record["camelot_duration_s"] = time.monotonic() - started
            except ImportError:
                record["camelot_skipped"] = "dependency missing"
            except Exception as exc:
                record["camelot_error"] = type(exc).__name__
            types = {"p1_t1": {1: "integer", 2: "currency", 3: "currency"}} if file_id == "synthetic_digital_001" else {}
            data = export_xlsx(document, numeric_columns=types)
            (directory / "result.xlsx").write_bytes(data)
            wb = load_workbook(BytesIO(data), data_only=False)
            expected_count = len(expected_tables)
            record["xlsx_table_count_matches"] = len(actual_tables) == expected_count
            exact = checked = content_ok = 0
            critical_ok = critical_checked = critical_content_ok = 0

            def content_matches(cell, expected: str, table_index: int, column: int, row_index: int) -> bool:
                if cell.data_type == "f":
                    return False
                if str(cell.value if cell.value is not None else "") == expected:
                    return True
                kind = types.get(f"p1_t{table_index}", {}).get(column)
                if kind and row_index > 0 and cell.data_type == "n":
                    try:
                        return Decimal(str(cell.value)) == Decimal(expected.strip().removeprefix("$").replace(",", ""))
                    except InvalidOperation:
                        return False
                return False
            for index, table in enumerate(expected_tables, 1):
                sheet_name = f"Table_{index}"
                if sheet_name not in wb:
                    checked += sum(len(values) for values in table["rows"])
                    critical_checked += len(table.get("critical_cells", []))
                    continue
                ws = wb[sheet_name]
                for r, values in enumerate(table["rows"], 1):
                    for c, expected in enumerate(values, 1):
                        checked += 1
                        cell = ws.cell(r, c)
                        if cell.data_type != "f" and str(cell.value if cell.value is not None else "") == expected:
                            exact += 1
                        content_ok += content_matches(cell, expected, index, c - 1, r - 1)
                for critical in table.get("critical_cells", []):
                    critical_checked += 1
                    cell = ws.cell(critical["row"] + 1, critical["col"] + 1)
                    critical_ok += cell.data_type != "f" and str(cell.value if cell.value is not None else "") == critical["value"]
                    critical_content_ok += content_matches(cell, critical["value"], index, critical["col"], critical["row"])
            # Typed synthetic cells are checked separately from verbatim text.
            record["xlsx_verbatim_cells"] = {"matched": exact, "checked": checked}
            record["xlsx_content_cells"] = {"matched": content_ok, "checked": checked}
            record["xlsx_critical_verbatim"] = {"matched": critical_ok, "checked": critical_checked}
            record["xlsx_critical_content"] = {"matched": critical_content_ok, "checked": critical_checked}
            if types:
                record["xlsx_typed_values_verified"] = wb["Table_1"]["B2"].value == 2 and wb["Table_1"]["D5"].value == 750.5
            original = wb["_Original"]
            expected_original = [c.content for t in document.get_all_tables() for c in t.cells]
            record["xlsx_original_text_preserved"] = [str(original.cell(r, 5).value or "") for r in range(2, original.max_row + 1)] == expected_original
            wb.close()
        if reference.get("scope") == "selected_pages_text" and row.category.startswith("scanned_"):
            record["ocr_full_page_results"] = []
            for page_reference in reference["pages"]:
                number = page_reference["page_number"]
                for psm in (3, 6):
                    started = time.monotonic()
                    result = {"page_number": number, "engine": "tesseract", "psm": psm, "dpi": 300,
                              "reference_characters": len(normalize(page_reference["text"]))}
                    try:
                        text = ocr_page_text(row.pdf_path, number, psm=psm)
                        result["cer_whitespace_normalized"] = character_error_rate(normalize(page_reference["text"]), normalize(text))
                        result["cer_raw"] = character_error_rate(page_reference["text"], text)
                        (directory / f"ocr_p{number}_psm{psm}.txt").write_text(text, encoding="utf-8")
                    except Exception as exc:
                        result["error"] = type(exc).__name__
                    result["duration_s"] = time.monotonic() - started
                    record["ocr_full_page_results"].append(result)
            record["ocr_scope_note"] = "Complete foreground text on selected pages only; not the full source document"
        if reference.get("scope") == "full_text":
            expected = normalize("\n".join(p.get("text", "") for p in reference["pages"]))
            actual = normalize(document.get_all_text())
            record["digital_text_cer_whitespace_normalized"] = character_error_rate(expected, actual)
            (directory / "result.txt").write_text(document.get_all_text(), encoding="utf-8")
            try:
                from docx import Document
                from packages.exporters.docx import export_docx
                data = export_docx(document)
                (directory / "result.docx").write_bytes(data)
                reopened = Document(BytesIO(data))
                content = []
                for paragraph in reopened.paragraphs:
                    if paragraph.text == "Conversion warnings":
                        break
                    content.append(paragraph.text)
                record["docx_text_cer_whitespace_normalized"] = character_error_rate(expected, normalize("\n".join(content)))
                reopened.add_paragraph("F0 editability verification")
                buffer = BytesIO()
                reopened.save(buffer)
                record["docx_edit_save_reopen"] = Document(BytesIO(buffer.getvalue())).paragraphs[-1].text == "F0 editability verification"
                try:
                    record["docx_alternative_route"] = pdfium_text_docx_baseline(
                        row.pdf_path, directory / "result_pdfium_baseline.docx", expected)
                except Exception as exc:
                    record["docx_alternative_error"] = type(exc).__name__
            except ImportError:
                record["docx_skipped"] = "dependency missing"
        regions = [(p["page_number"], region) for p in reference["pages"] for region in p.get("text_regions", [])]
        if regions:
            try:
                import pypdfium2
                import pytesseract
                # Reuse the repository's Windows binary configuration.
                from benchmarks import ocr_benchmark  # noqa: F401
                with pypdfium2.PdfDocument(str(row.pdf_path)) as pdf:
                    scores = []
                    for number, region in regions:
                        page = pdf[number - 1]
                        try:
                            bitmap = page.render(scale=300 / 72)
                            image = bitmap.to_pil()
                            x0, y0, x1, y1 = region["bbox_relative"]
                            crop = image.crop((round(x0 * image.width), round(y0 * image.height),
                                               round(x1 * image.width), round(y1 * image.height)))
                            actual = pytesseract.image_to_string(crop, lang="spa+eng", config="--psm 6")
                            scores.append(character_error_rate(normalize(region["text"]), normalize(actual)))
                            crop.close()
                            image.close()
                            bitmap.close()
                        finally:
                            page.close()
                    record["ocr_region_cer_whitespace_normalized"] = scores
                    record["ocr_scope_note"] = "Cropped region only; not whole-page CER or scanned_clean gate evidence"
            except Exception as exc:
                record["ocr_region_error"] = f"{type(exc).__name__}: {exc}"
        results.append(record)
    packages = {}
    for package in ["pdfplumber", "openpyxl", "pydantic", "python-docx", "camelot-py"]:
        try:
            packages[package] = version(package)
        except Exception:
            packages[package] = "not installed"
    report = {"gate_f0": "PENDING: small sample, partial document references; structure, costs and licenses not closed",
              "versions": packages, "results": results}
    (output / "report.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(run(), indent=2, ensure_ascii=False))
