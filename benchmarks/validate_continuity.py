"""Validate visually annotated boundary rows independently of whole-page quality."""
from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook

from benchmarks.common import load_ground_truth, load_manifest
from packages.converters.digital import extract_digital
from packages.exporters.xlsx import export_xlsx


def run() -> dict:
    row = next(row for row in load_manifest() if row.file_id == "estado_cuenta_004")
    if not row.authorized:
        raise ValueError("Source not authorized")
    reference = load_ground_truth(row.file_id)
    case = reference["continuity_cases"][0]
    document = extract_digital(row.pdf_path, page_numbers=[18, 19])
    tables = document.get_all_tables()
    first, second = tables
    expected_rows = case["boundary_rows"]
    matched = 0
    for expected in expected_rows:
        table = next(table for table in tables if table.page_number == expected["page_number"])
        actual = [next(cell.content for cell in table.cells if cell.row_index == expected["row"] and cell.col_index == col)
                  for col in range(table.columns)]
        matched += sum(value == truth for value, truth in zip(actual, expected["values"]))
    data = export_xlsx(document)
    workbook = load_workbook(BytesIO(data), data_only=False)
    preserved = all(str(workbook[f"Table_{i}"].cell(cell.row_index + 1, cell.col_index + 1).value or "") == cell.content
                    for i, table in enumerate(tables, 1) for cell in table.cells)
    workbook.close()
    output = Path(".generated/f0/continuity")
    output.mkdir(parents=True, exist_ok=True)
    (output / "result.xlsx").write_bytes(data)
    (output / "document.ir.json").write_text(document.model_dump_json(indent=2), encoding="utf-8")
    report = {"file_id": row.file_id, "source_pages": document.page_count, "evaluated_pages": [18, 19],
              "continuation_candidate_detected": second.continuation_of == first.id,
              "boundary_cells": {"matched": matched, "checked": sum(len(r["values"]) for r in expected_rows)},
              "source_segment_rows": [t.rows for t in tables], "all_extracted_cells_preserved_in_xlsx": preserved,
              "scope": "boundary rows only; split-row reconstruction and full page 18 truth not validated"}
    (output / "report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(run(), indent=2))
