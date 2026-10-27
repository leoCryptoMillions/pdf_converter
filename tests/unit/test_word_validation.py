from io import BytesIO

from docx import Document

from benchmarks.validate_word import content_comparison, word_text


def test_comparison_detects_duplicates_and_does_not_normalize_case():
    assert content_comparison('A B', 'A B B')['extra_characters'] == 1
    assert content_comparison('Ab', 'ab')['missing_characters'] == 1
    assert content_comparison('A\nB', 'AB')['exact_character_multiset']
    reordered = content_comparison('AB', 'BA')
    assert reordered['exact_character_multiset']
    assert 'does not verify reading order' in reordered['scope']


def test_word_text_visits_merged_cells_only_once():
    word = Document()
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Only once'
    buffer = BytesIO()
    word.save(buffer)
    assert word_text(Document(BytesIO(buffer.getvalue()))).count('Only once') == 1
