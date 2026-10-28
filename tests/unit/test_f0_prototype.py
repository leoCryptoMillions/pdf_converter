from io import BytesIO
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook
import pytest
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from benchmarks.common import character_error_rate, table_f1
from packages.converters.digital import PDFValidationError, extract_digital
from packages.exporters.xlsx import export_xlsx


@pytest.fixture
def digital_pdf(tmp_path):
    path = tmp_path / "table.pdf"
    rows = [["Identifier", "Quantity", "Amount", "Notes"],
            ["000123456789012345", "2", "$350.50", '=HYPERLINK("https://example.invalid")'],
            ["000123456789012345", "2", "$350.50", '=HYPERLINK("https://example.invalid")'],
            ["ABC", "012", "1.234,56", "@SUM(A1)"]]
    table = Table(rows, colWidths=[115, 55, 70, 240])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                               ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    SimpleDocTemplate(str(path)).build([table])
    return path, rows


@pytest.mark.mvp
def test_real_extraction_export_preserves_duplicates_ids_and_formulas(digital_pdf):
    source, expected = digital_pdf
    document = extract_digital(source)
    assert document.pages[0].text_coverage is None
    assert document.text_extraction_confidence is None
    assert document.has_warnings()
    table = document.get_all_tables()[0]
    assert table.rows == 4
    assert table.confidence is None
    wb = load_workbook(BytesIO(export_xlsx(document)), data_only=False)
    ws = wb["Table_1"]
    for r, row in enumerate(expected, 1):
        for c, text in enumerate(row, 1):
            assert ws.cell(r, c).value == text
            assert ws.cell(r, c).data_type == "s"
    assert len(wb["_Original"]["E"]) == 17
    assert {"_Origin", "_Original", "_Warnings"} <= set(wb.sheetnames)
    wb.close()


@pytest.mark.mvp
def test_numeric_types_are_explicit_and_ambiguous_values_stay_text(digital_pdf):
    document = extract_digital(digital_pdf[0])
    data = export_xlsx(document, numeric_columns={"p1_t1": {1: "integer", 2: "currency"}})
    wb = load_workbook(BytesIO(data), data_only=False)
    ws = wb["Table_1"]
    assert ws["A2"].value == "000123456789012345"
    assert ws["B2"].value == 2
    assert ws["C2"].value == 350.5
    assert ws["B4"].value == "012"
    assert ws["C4"].value == "1.234,56"
    assert wb["_Original"]["E8"].value == "$350.50"
    warnings = list(wb["_Warnings"].values)
    assert sum(row[0] == "NUMERIC_VALUE_KEPT_AS_TEXT" for row in warnings) == 2
    wb.close()


@pytest.mark.mvp
def test_excel_precision_guard(digital_pdf):
    document = extract_digital(digital_pdf[0])
    document.pages[0].tables[0].cells[5].content = "1234567890123456"
    wb = load_workbook(BytesIO(export_xlsx(document, numeric_columns={"p1_t1": {1: "integer"}})))
    assert wb["Table_1"]["B2"].value == "1234567890123456"
    wb.close()


@pytest.mark.mvp
def test_xlsx_keeps_full_page_text_even_when_tables_are_found(tmp_path):
    # Real-world bug (2026-10-27, a real vehicle-repair quotation): export_xlsx
    # only ever wrote detected tables, silently dropping everything outside
    # their bbox (company letterhead, client name, serial/model, footer
    # note) whenever ANY table was found - a document with real tables came
    # out looking nothing like the original.
    from reportlab.pdfgen import canvas

    path = tmp_path / "quotation.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(400, 400))
    pdf.setFont("Helvetica", 10)
    pdf.drawString(40, 370, "ACME TALLER")
    pdf.drawString(40, 355, "CLIENTE: Jane Doe")
    table = Table([["CANT.", "DESCRIPCION"], ["1", "FILTRO"]], colWidths=[50, 100])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    table.wrapOn(pdf, 150, 40)
    table.drawOn(pdf, 40, 280)
    pdf.drawString(40, 260, "Precios sujetos a cambio sin previo aviso.")
    pdf.save()

    document = extract_digital(path)
    assert document.get_all_tables()  # a table was in fact detected
    wb = load_workbook(BytesIO(export_xlsx(document)), data_only=False)
    assert wb.sheetnames[0] == "Texto completo"
    full_text = "\n".join(str(c.value) for row in wb["Texto completo"].iter_rows() for c in row if c.value)
    assert "ACME TALLER" in full_text
    assert "CLIENTE: Jane Doe" in full_text
    assert "Precios sujetos a cambio sin previo aviso." in full_text
    wb.close()


@pytest.mark.mvp
def test_xlsx_flags_a_cell_that_may_span_multiple_source_rows(tmp_path):
    # Real-world bug (2026-10-27, same quotation): a SUBTOTAL/IVA/TOTAL block
    # had no row-dividing line under its label column (only under its value
    # column), so pdfplumber's grid merged three source rows' labels into one
    # cell, joined by newlines, while the values stayed correctly split.
    from reportlab.platypus import Paragraph
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    path = tmp_path / "wrapped-cell.pdf"
    rows = [["Header A", "Header B"],
            [Paragraph("SUBTOTAL<br/>IVA<br/>TOTAL", styles["Normal"]), "100"]]
    table = Table(rows, colWidths=[120, 50])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(str(path)).build([table])

    document = extract_digital(path)
    cell = next(c for c in document.get_all_tables()[0].cells if "\n" in c.content)
    assert cell.content.count("\n") >= 1
    assert any(w.code == "MULTI_LINE_CELL_MAY_SPAN_ROWS" for w in cell.warnings)


@pytest.mark.mvp
def test_xlsx_visual_page_places_same_row_blocks_without_overlap(tmp_path):
    # Design-review correction: two free-text blocks sitting side by side on
    # the same visual row (e.g. a real quotation's "SERIE: ..." and
    # "MODELO ..." on one line) must each stop their column span at the
    # other's left edge rather than both reaching for the page's rightmost
    # column, or their merge_cells ranges collide - verified to actually
    # raise AttributeError ("MergedCell ... read-only") without this fix.
    from reportlab.pdfgen import canvas

    path = tmp_path / "side_by_side.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(400, 200))
    pdf.setFont("Helvetica", 10)
    pdf.drawString(40, 150, "SERIE: P7401430")
    pdf.drawString(220, 150, "MODELO E10x")
    pdf.drawString(350, 150, "FIN")
    pdf.drawString(40, 100, "Mas contenido debajo de la linea")
    pdf.save()

    document = extract_digital(path)
    page = document.pages[0]
    assert len(page.blocks) == 4

    data = export_xlsx(document)  # Must not raise.
    wb = load_workbook(BytesIO(data), data_only=False)
    ws = wb[f"Pagina_{page.page_number}"]

    from packages.exporters.xlsx import _build_page_grid, _locate

    xs, ys = _build_page_grid(page)
    serie = next(b for b in page.blocks if b.content.startswith("SERIE"))
    modelo = next(b for b in page.blocks if b.content.startswith("MODELO"))
    fin = next(b for b in page.blocks if b.content == "FIN")
    serie_col, modelo_col, fin_col = (_locate(xs, b.location.x0) for b in (serie, modelo, fin))
    assert serie_col < modelo_col < fin_col  # Same row, strictly increasing columns.
    shared_row = _locate(ys, serie.location.y0)
    assert shared_row == _locate(ys, modelo.location.y0) == _locate(ys, fin.location.y0)

    from openpyxl.utils.cell import range_boundaries

    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if min_row - 1 != shared_row:
            continue  # A different element's (e.g. the line below's) own row.
        # No same-row merge may reach into or past a later same-row neighbour's column.
        if min_col - 1 == serie_col:
            assert max_col - 1 < modelo_col
        if min_col - 1 == modelo_col:
            assert max_col - 1 < fin_col
    wb.close()


@pytest.mark.mvp
def test_xlsx_captures_and_renders_a_colored_highlight_band(tmp_path):
    # Real-world case this feature targets: a green SUBTOTAL/IVA/TOTAL band
    # (CLAUDE.md's "highlight cells requiring review" idea, but sourced from
    # an actually observed PDF fill rect, never a heuristic).
    from reportlab.lib import colors as rl_colors

    path = tmp_path / "highlight.pdf"
    rows = [["Label", "Value"], ["SUBTOTAL", "100.00"]]
    table = Table(rows, colWidths=[100, 60])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, rl_colors.black),
        ("BACKGROUND", (0, 1), (-1, 1), rl_colors.HexColor("#00B050")),
    ]))
    SimpleDocTemplate(str(path)).build([table])

    document = extract_digital(path)
    table_ir = document.get_all_tables()[0]
    header_cells = [c for c in table_ir.cells if c.row_index == 0]
    highlighted = [c for c in table_ir.cells if c.row_index == 1]
    assert all(c.fill_color is None for c in header_cells)
    assert highlighted and all(c.fill_color == "00B050" for c in highlighted)

    from packages.exporters.xlsx import _build_page_grid, _locate

    wb = load_workbook(BytesIO(export_xlsx(document)))
    page = document.pages[0]
    ws = wb[f"Pagina_{page.page_number}"]
    xs, ys = _build_page_grid(page)
    row0 = _locate(ys, table_ir.location.y0)
    col0 = _locate(xs, table_ir.location.x0)
    cell = next(c for c in highlighted if c.col_index == 0)
    target = ws.cell(row0 + cell.row_index + 1, col0 + cell.col_index + 1)
    assert target.fill.fgColor.rgb.upper().endswith("00B050")
    wb.close()


@pytest.mark.mvp
def test_xlsx_visual_page_grows_rows_so_wrapped_text_is_not_clipped(tmp_path):
    # Real-world bug (2026-10-27, found rendering real documents through
    # actual Excel, not just openpyxl round-trips): Pagina_N row heights were
    # taken straight from the PDF's own geometry, which only ever sized a row
    # for ONE line. A narrow Excel column (its width is an approximation of
    # the PDF point width, see known limitation) wraps that same text into
    # several lines, and a one-line-tall row then shows only a sliver of the
    # first wrapped line - e.g. "REFACCIONES" rendered as "REF". The fix must
    # grow the row to fit every wrapped line, not just preserve PDF geometry.
    from reportlab.pdfgen import canvas

    path = tmp_path / "narrow_column.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(200, 150))
    pdf.setFont("Helvetica", 10)
    # A single narrow-looking line of text that, once placed in export_xlsx's
    # approximated (narrow) column width, needs several wrapped lines.
    pdf.drawString(20, 120, "REFACCIONES Y REPUESTOS AUTOMOTRICES COMPLETOS")
    pdf.drawString(20, 60, "Linea de abajo")
    pdf.save()

    document = extract_digital(path)
    page = document.pages[0]
    block = next(b for b in page.blocks if b.content.startswith("REFACCIONES"))

    from packages.exporters.xlsx import (
        _build_page_grid, _col_width, _estimate_wrapped_lines, _locate,
    )

    xs, ys = _build_page_grid(page)
    row_start = _locate(ys, block.location.y0)
    col_start = _locate(xs, block.location.x0)
    avail_chars = _col_width(xs, col_start)  # The block's own single column; no later neighbour on its row.
    expected_lines = _estimate_wrapped_lines(block.content, avail_chars)
    assert expected_lines > 1  # The scenario only tests something if Excel actually has to wrap it.

    data = export_xlsx(document)
    wb = load_workbook(BytesIO(data))
    ws = wb[f"Pagina_{page.page_number}"]
    height = ws.row_dimensions[row_start + 1].height
    single_line_height = 10.0 * 1.2  # font_size * line-height multiplier, see _required_row_height.
    # A one-line-tall row would clip everything past the first wrapped line;
    # the fix must size the row for every wrapped line the text actually needs.
    assert height >= expected_lines * single_line_height
    wb.close()


@pytest.mark.mvp
def test_xlsx_embeds_captured_images_and_discloses_undiscloseable_ones(tmp_path):
    # Real-world bug (2026-10-27, a real SAT "Constancia de Situacion Fiscal"):
    # extract_digital captures embedded images (logos, QR codes, barcodes) in
    # page.images, but export_xlsx's Pagina_N overlay never read page.images
    # at all - every image silently vanished from the visual sheet with no
    # warning, hiding real degradation. 'graphics' kind (a full-page vector
    # re-render) is deliberately NOT embedded (it would float on top of and
    # obscure cell text/fills) but must still be disclosed, never silent.
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image as PILImage

    png_path = tmp_path / "logo.png"
    PILImage.new("RGB", (40, 20), color=(200, 30, 30)).save(png_path)

    path = tmp_path / "with_image.pdf"
    pdf = canvas.Canvas(str(path), pagesize=(300, 200))
    pdf.setFont("Helvetica", 10)
    pdf.drawString(20, 170, "HEADER TEXT")
    pdf.drawImage(ImageReader(str(png_path)), 20, 100, width=40, height=20)
    table = Table([["A", "B"], ["1", "2"]], colWidths=[40, 40])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    table.wrapOn(pdf, 100, 50)
    table.drawOn(pdf, 20, 40)
    pdf.save()

    document = extract_digital(path)
    page = document.pages[0]
    assert any(img.kind == "embedded" for img in page.images)

    data = export_xlsx(document)
    wb = load_workbook(BytesIO(data))
    ws = wb[f"Pagina_{page.page_number}"]
    assert len(ws._images) >= 1  # The embedded logo must actually be drawn into the sheet.

    warnings = list(wb["_Warnings"].values)
    image_warnings = [row for row in warnings if row[0] == "IMAGE_NOT_EMBEDDED"]
    embedded_ids = {img.id for img in page.images if img.kind == "embedded"}
    # An embedded image that was in fact drawn must never also be reported as
    # not embedded; a 'graphics' image that was deliberately skipped must be.
    assert not any(any(image_id in (row[2] or "") for image_id in embedded_ids) for row in image_warnings)
    if any(img.kind == "graphics" for img in page.images):
        assert image_warnings
    wb.close()


@pytest.mark.mvp
def test_dense_table_layout_enrichment_is_skipped_within_time_budget(tmp_path):
    # Safety regression surfaced by adversarial testing (2026-10-27): a
    # plausibly realistic dense ruled table (not an extreme input) drove
    # extract_digital's per-cell layout enrichment (within_bbox +
    # positioned_lines + cell_runs + detect_fill, each called independently
    # per cell with no spatial index) past apps/conversion_isolation.py's 90s
    # hard production timeout, so export_xlsx's own graceful
    # VISUAL_GRID_TOO_DENSE degrade was never reached - the user saw an
    # opaque timeout kill instead of a warning-bearing result. Skipping the
    # expensive per-cell enrichment above a cell-count budget (content and
    # row/col span stay correct) keeps extraction inside the production
    # timeout with a wide safety margin, so the degrade is actually reachable.
    import time

    from packages.converters.digital import LAYOUT_ENRICHMENT_CELL_LIMIT

    n = 58  # Comfortably above LAYOUT_ENRICHMENT_CELL_LIMIT once pdfplumber lays the grid out.
    path = tmp_path / "dense_grid.pdf"
    rows = [[f"{r}-{c}" for c in range(n)] for r in range(n)]
    table = Table(rows, colWidths=[20] * n, rowHeights=[12] * n)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.black),
                               ("FONTSIZE", (0, 0), (-1, -1), 6)]))
    SimpleDocTemplate(str(path), pagesize=(n * 20 + 40, n * 12 + 40)).build([table])

    started = time.monotonic()
    document = extract_digital(path)
    elapsed = time.monotonic() - started

    table_ir = document.get_all_tables()[0]
    assert len(table_ir.cells) > LAYOUT_ENRICHMENT_CELL_LIMIT
    assert any(w.code == "TABLE_LAYOUT_ENRICHMENT_SKIPPED" for w in table_ir.warnings)
    # Content extraction (not the skipped enrichment) must remain correct.
    assert any(c.content == "0-0" for c in table_ir.cells)
    assert all(c.row_span >= 1 and c.col_span >= 1 for c in table_ir.cells)
    # Generous bound: well under the 90s production timeout (apps/
    # conversion_isolation.py), with ample margin for slower CI machines -
    # before this fix, a table this size measured close to 100s.
    assert elapsed < 45

    data = export_xlsx(document)  # Must still complete without raising or hanging.
    assert data


@pytest.mark.mvp
def test_metadata_does_not_become_a_formula(digital_pdf):
    document = extract_digital(digital_pdf[0])
    document.engines_used[0].engine_version = "=1+1"
    document.pages[0].tables[0].id = "=2+2"
    wb = load_workbook(BytesIO(export_xlsx(document)), data_only=False)
    assert all(cell.data_type != "f" for sheet in wb for row in sheet for cell in row)
    wb.close()


@pytest.mark.parametrize("text", ["x" * 32768, "bad\x00text"], ids=["too_long", "control_character"])
@pytest.mark.mvp
def test_unrepresentable_excel_text_fails_instead_of_truncating(digital_pdf, text):
    document = extract_digital(digital_pdf[0])
    document.pages[0].tables[0].cells[0].content = text
    with pytest.raises(ValueError, match="losslessly"):
        export_xlsx(document)


@pytest.mark.mvp
def test_duplicate_cell_coordinates_rejected(digital_pdf):
    document = extract_digital(digital_pdf[0])
    table = document.pages[0].tables[0]
    table.cells.append(table.cells[0])
    with pytest.raises(ValueError, match="duplicate"):
        export_xlsx(document)


@pytest.mark.mvp
def test_docx_contains_editable_table_without_duplicate_text(digital_pdf):
    from docx import Document
    from packages.exporters.docx import export_docx

    document = extract_digital(digital_pdf[0])
    reopened = Document(BytesIO(export_docx(document)))
    assert len(reopened.tables) == 1
    assert reopened.tables[0].cell(1, 0).text == "000123456789012345"
    assert all("Identifier" not in paragraph.text for paragraph in reopened.paragraphs)
    reopened.tables[0].cell(1, 1).text = "Edited"
    output = BytesIO()
    reopened.save(output)
    assert Document(BytesIO(output.getvalue())).tables[0].cell(1, 1).text == "Edited"


@pytest.mark.mvp
def test_input_limits_and_signature(digital_pdf, tmp_path):
    with pytest.raises(PDFValidationError, match="size"):
        extract_digital(digital_pdf[0], max_bytes=10)
    with pytest.raises(PDFValidationError, match="page limit"):
        extract_digital(digital_pdf[0], max_pages=0)
    invalid = tmp_path / "invalid.pdf"
    invalid.write_bytes(b"not a pdf")
    with pytest.raises(PDFValidationError, match="signature"):
        extract_digital(invalid)
    invalid.write_bytes(b"%PDF-1.7\nbroken")
    with pytest.raises(PDFValidationError, match="parsed"):
        extract_digital(invalid)


@pytest.mark.mvp
def test_encrypted_pdf_rejected(digital_pdf, tmp_path):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(digital_pdf[0]))
    writer.encrypt("secret")
    path = tmp_path / "encrypted.pdf"
    writer.write(path)
    with pytest.raises(PDFValidationError):
        extract_digital(path)


@pytest.mark.mvp
def test_cli_outputs_and_refuses_overwrite(digital_pdf, tmp_path):
    output = tmp_path / "conversion"
    command = [sys.executable, "-m", "apps.convert", str(digital_pdf[0]), "--output", str(output),
               "--formats", "xlsx", "docx", "txt"]
    result = subprocess.run(command, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert (output / "result.xlsx").is_file()
    assert (output / "result.docx").is_file()
    assert "Identifier" in (output / "result.txt").read_text(encoding="utf-8")
    assert (output / "summary.json").is_file()
    assert subprocess.run(command, capture_output=True).returncode == 1


@pytest.mark.parametrize("expected,actual,distance", [
    ("abc", "abc", 0), ("kitten", "sitting", 3), ("abab", "baba", 2),
    ("abc", "", 3), ("a", "bbbb", 4),
])
def test_exact_cer(expected, actual, distance):
    assert character_error_rate(expected, actual) == pytest.approx(distance / len(expected))


def test_table_metric_detects_missing_duplicate():
    score = table_f1([{"rows": [["A"], ["A"]]}], [[["A"]]])
    assert score["recall"] == 0.5
    assert score["f1"] == pytest.approx(2 / 3)


def test_draft_reference_is_not_scored(tmp_path, monkeypatch):
    import benchmarks.common as common

    (tmp_path / "draft.json").write_text('{"review_status": "draft"}', encoding="utf-8")
    monkeypatch.setattr(common, "GROUND_TRUTH_DIR", tmp_path)
    assert common.load_ground_truth("draft") is None


def test_reference_hash_mismatch_is_rejected(tmp_path, monkeypatch):
    import benchmarks.common as common

    (tmp_path / "local").mkdir()
    (tmp_path / "local" / "test.json").write_text(
        '{"review_status": "verified", "source_sha256": "wrong_hash"}', encoding="utf-8")
    (tmp_path / "source.pdf").write_bytes(b"different source")
    monkeypatch.setattr(common, "GROUND_TRUTH_DIR", tmp_path)
    monkeypatch.setattr(common, "CORPUS_DIR", tmp_path)
    row = common.ManifestRow("test", "source.pdf", "digital_simple_table", 1, "synthetic", True, True)
    monkeypatch.setattr(common, "load_manifest", lambda: [row])
    with pytest.raises(ValueError, match="hash mismatch"):
        common.load_ground_truth("test")


def test_reference_path_traversal_rejected():
    from benchmarks.common import load_ground_truth

    with pytest.raises(ValueError, match="Invalid reference ID"):
        load_ground_truth("../reference")


@pytest.mark.parametrize("pages", [[], [0], [-1], [2], [1, 1], [True]],
                         ids=["empty", "zero", "negative", "outside", "duplicate", "boolean"])
@pytest.mark.mvp
def test_invalid_selected_pages_rejected(digital_pdf, pages):
    with pytest.raises(PDFValidationError, match="selected pages"):
        extract_digital(digital_pdf[0], page_numbers=pages)


@pytest.mark.mvp
def test_page_selection_retains_original_page_number_and_source_count(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "two_pages.pdf"
    writer = canvas.Canvas(str(path))
    writer.drawString(72, 700, "First page")
    writer.showPage()
    writer.drawString(72, 700, "Second page")
    writer.save()
    document = extract_digital(path, page_numbers=[2])
    assert document.page_count == 2
    assert len(document.pages) == 1
    assert document.pages[0].page_number == 2
    assert document.pages[0].extracted_text == "Second page"
    assert document.extraction_options["selected_pages"] == [2]
    assert "First page" not in document.get_all_text()


def test_scoped_table_benchmark_passes_only_annotated_pages(digital_pdf, monkeypatch):
    import benchmarks.tables_benchmark as benchmark
    from benchmarks.common import ManifestRow

    row = ManifestRow("scoped", digital_pdf[0].name, "digital_complex_table", 19, "synthetic", True, True)
    monkeypatch.setattr(benchmark, "load_manifest", lambda: [row])
    monkeypatch.setattr(benchmark, "usable_rows", lambda rows, categories: (rows, []))
    monkeypatch.setattr(benchmark, "load_ground_truth", lambda ident: {
        "scope": "selected_pages_tables", "pages": [{"page_number": 19, "tables": [{"rows": [["Value"]]}]}]})
    received = []

    def extractor(path, page_numbers=None):
        received.append(page_numbers)
        return [[["Value"]]]

    monkeypatch.setattr(benchmark, "ENGINES", {"test": extractor})
    results, skipped = benchmark.run()
    assert received == [[19]]
    assert results[0].metrics["f1"] == 1
    assert results[0].metrics["reference_scope"] == "selected_pages_tables"


def test_full_document_ocr_does_not_score_partial_page_reference(monkeypatch):
    import benchmarks.ocr_benchmark as benchmark
    from benchmarks.common import ManifestRow

    row = ManifestRow("scoped", "source.pdf", "scanned_clean", 14, "synthetic", True, True)
    monkeypatch.setattr(benchmark, "load_manifest", lambda: [row])
    monkeypatch.setattr(benchmark, "usable_rows", lambda rows, categories: (rows, []))
    monkeypatch.setattr(benchmark, "load_ground_truth", lambda ident: {
        "scope": "selected_pages_text", "pages": [{"page_number": 1, "text": "Cover"}]})
    monkeypatch.setattr(benchmark, "ENGINES", {"test": lambda path: "Cover"})
    results, skipped = benchmark.run()
    assert results[0].metrics["cer"] is None
    assert "partial-document" in results[0].metrics["note"]


@pytest.mark.mvp
def test_ocr_pixel_limit_rejects_page_before_rasterization(digital_pdf):
    from benchmarks.validate_prototype import ocr_page_text

    with pytest.raises(ValueError, match="pixel budget"):
        ocr_page_text(digital_pdf[0], 1, max_pixels=1)


def test_independent_pdfium_word_baseline_preserves_editable_text(tmp_path):
    from reportlab.pdfgen import canvas
    from benchmarks.validate_prototype import pdfium_text_docx_baseline

    path = tmp_path / "text.pdf"
    writer = canvas.Canvas(str(path))
    writer.drawString(72, 700, "Expected text")
    writer.save()
    result = pdfium_text_docx_baseline(path, tmp_path / "baseline.docx", "Expected text")
    assert result["cer_whitespace_normalized"] == 0
    assert result["edit_save_reopen"] is True


@pytest.fixture
def segmented_document():
    from packages.document_ir.models import Cell, DocumentIR, Page, Rectangle, Table

    pages = []
    for number, top in [(1, 700), (2, 70)]:
        cells = [Cell(content=f"p{number}_{r}_{c}", row_index=r, col_index=c,
                      extraction_method="direct", original_source=Rectangle(
                          x0=72 + c * 100, x1=172 + c * 100, y0=top + r * 30, y1=top + (r + 1) * 30))
                 for r in range(2) for c in range(2)]
        table = Table(id=f"p{number}_t1", page_number=number, rows=2, columns=2,
                      cells=cells, extraction_method="direct", source_engine="test",
                      location=Rectangle(x0=72, x1=272, y0=top, y1=top + 60))
        pages.append(Page(page_number=number, width=612, height=792, classification="digital",
                          image_coverage=0, tables=[table]))
    return DocumentIR(document_id="test", source_filename="test.pdf", source_hash="test",
                      page_count=2, pages=pages, engines_used=[])


def test_continuity_links_preserve_rows_and_warn_idempotently(segmented_document):
    from packages.converters.continuity import link_table_continuations

    tables = segmented_document.get_all_tables()
    original = [c.content for t in tables for c in t.cells]
    link_table_continuations(segmented_document)
    link_table_continuations(segmented_document)
    assert tables[1].continuation_of == tables[0].id
    assert len(tables[1].warnings) == 2
    assert [c.content for t in tables for c in t.cells] == original
    wb = load_workbook(BytesIO(export_xlsx(segmented_document)))
    assert wb["Table_1"].max_row == 2
    assert wb["Table_2"].max_row == 2
    assert any("continuation candidate" in str(row[0]) for row in wb["_Origin"].values)
    wb.close()


@pytest.mark.parametrize("mismatch", ["gap", "geometry", "rotation", "edge", "span"])
def test_continuity_rejects_insufficient_evidence(segmented_document, mismatch):
    from packages.converters.continuity import link_table_continuations

    previous, current = segmented_document.pages
    if mismatch == "gap":
        current.page_number = 3
    elif mismatch == "geometry":
        for cell in current.tables[0].cells:
            cell.original_source.x0 += 15
            cell.original_source.x1 += 15
    elif mismatch == "rotation":
        current.rotation = 90
    elif mismatch == "edge":
        previous.tables[0].location.y1 = 400
    else:
        current.tables[0].cells[0].col_span = 2
    link_table_continuations(segmented_document)
    assert current.tables[0].continuation_of is None


def test_cost_projection_requires_actual_tariffs():
    from benchmarks.resource_benchmark import cost_scenarios

    unknown = cost_scenarios(1, pages=[3600])[0]
    assert unknown["compute_cost_estimate"] is None
    assert unknown["total_estimate"] is None
    priced = cost_scenarios(1, pages=[3600], hourly_rate=3, fixed_monthly=4)[0]
    assert priced["sequential_worker_hours"] == 1
    assert priced["total_estimate"] == 7


@pytest.mark.parametrize("seconds", [-1, float("nan"), float("inf")])
def test_cost_projection_rejects_invalid_resource_measurements(seconds):
    from benchmarks.resource_benchmark import cost_scenarios

    with pytest.raises(ValueError):
        cost_scenarios(seconds)
