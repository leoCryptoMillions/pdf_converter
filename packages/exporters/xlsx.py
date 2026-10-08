"""Loss-conscious XLSX export. Original values always remain text."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

from packages.document_ir.models import DocumentIR


def put_text(cell, value: str) -> None:
    if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
        raise ValueError("Text cannot be represented losslessly in an XLSX cell")
    cell.value = value
    cell.data_type = "s"  # Includes =, +, -, @ content; never create a formula.
    cell.number_format = "@"


def export_xlsx(document: DocumentIR, *, numeric_columns: dict[str, dict[int, str]] | None = None) -> bytes:
    """Numeric conversion requires an explicit table/column map and en_US syntax.

    Example: {'p1_t1': {1: 'integer', 2: 'currency', 3: 'currency'}}.
    Dates and identifiers stay text. Excel's 15-digit precision is enforced.
    """
    numeric_columns = numeric_columns or {}
    wb = Workbook()
    wb.remove(wb.active)
    original = wb.create_sheet("_Original")
    original.append(["Table", "Page", "Row", "Column", "Original text"])
    issues = []
    for i, table in enumerate(document.get_all_tables(), 1):
        ws = wb.create_sheet(f"Table_{i}")
        if table.rows > 1048576 or table.columns > 16384:
            raise ValueError("Table exceeds Excel dimensions")
        occupied = set()
        for source in table.cells:
            r, c = source.row_index, source.col_index
            if not (0 <= r < table.rows and 0 <= c < table.columns) or (r, c) in occupied:
                raise ValueError("Invalid or duplicate cell coordinates")
            occupied.add((r, c))
            target = ws.cell(r + 1, c + 1)
            put_text(target, source.content)
            original.append([table.id, table.page_number, r, c, None])
            put_text(original.cell(original.max_row, 5), source.content)
            kind = numeric_columns.get(table.id, {}).get(c)
            if kind and r > 0 and source.content.strip():
                if kind not in {"integer", "decimal", "currency"}:
                    raise ValueError("Unsupported numeric type")
                cleaned = source.content.strip()
                if kind == "currency":
                    cleaned = cleaned.removeprefix("$").strip()
                pattern = r"-?(?:0|[1-9]\d*|[1-9]\d{0,2}(?:,\d{3})+)(?:\.\d+)?"
                try:
                    if not re.fullmatch(pattern, cleaned):
                        raise ValueError()
                    value = Decimal(cleaned.replace(",", ""))
                    if (len(value.as_tuple().digits) > 15 or
                        (value != 0 and not -307 <= value.adjusted() <= 307) or
                        (kind == "currency" and value != value.quantize(Decimal("0.01"))) or
                        (kind == "integer" and value != value.to_integral_value())):
                        raise ValueError()
                    target.value = int(value) if kind == "integer" else value
                    target.number_format = "0" if kind == "integer" else '#,##0.00' if kind == "currency" else "General"
                except (ValueError, InvalidOperation):
                    issues.append(("NUMERIC_VALUE_KEPT_AS_TEXT", table.page_number,
                                   f"{table.id} R{r + 1} C{c + 1}: ambiguous format or precision; original retained."))
            target.alignment = Alignment(wrap_text=True, vertical="top")
            if table.has_header and r < table.header_rows:
                target.font = Font(bold=True)
            for issue in source.warnings:
                issues.append((issue.code, issue.page_number, issue.message))
        ws.freeze_panes = "A2" if table.has_header else None
        for issue in table.warnings:
            issues.append((issue.code, issue.page_number, issue.message))
    if not document.get_all_tables():
        ws = wb.create_sheet("Text")
        for page in document.pages:
            for line in page.extracted_text.splitlines():
                put_text(ws.cell(ws.max_row + 1 if ws.cell(1, 1).value is not None else 1, 1), line)
        issues.append(("NO_TABLES", 1, "No ruled tables detected; only extracted text is exported."))
    origin = wb.create_sheet("_Origin")
    for key, value in [("Source SHA256", document.source_hash), ("Schema", document.schema_version),
                       ("Source page count", str(document.page_count)),
                       ("Exported source pages", ", ".join(str(p.page_number) for p in document.pages)),
                       ("Numeric syntax", "en_US; explicit columns only"),
                       ("Quality", "unverified; needs_review")]:
        origin.append([key, None])
        put_text(origin.cell(origin.max_row, 2), value)
    for engine in document.engines_used:
        origin.append([engine.engine_name, engine.engine_version])
    for table in document.get_all_tables():
        origin.append([table.id, f"page {table.page_number}; bbox {table.location.model_dump()}"])
        if table.continuation_of:
            origin.append([f"{table.id} continuation candidate", table.continuation_of])
    for issue in document.warnings + [w for p in document.pages for w in p.warnings]:
        issues.append((issue.code, issue.page_number, issue.message))
    ws = wb.create_sheet("_Warnings")
    ws.append(["Code", "Page", "Message"])
    for code, page, message in issues:
        ws.append([None, page, None])
        put_text(ws.cell(ws.max_row, 1), code)
        put_text(ws.cell(ws.max_row, 3), message)
    result = BytesIO()
    # Metadata in an externally supplied IR is also untrusted text.
    for sheet in wb:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    put_text(cell, cell.value)
    wb.save(result)
    return result.getvalue()
