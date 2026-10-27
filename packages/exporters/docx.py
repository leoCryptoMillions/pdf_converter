"""Editable Word with measured PDF typography and page-relative layout.

Paragraph frames, floating tables and drawing anchors use standard OOXML.
See docs/adr/0005-word-layout.md for scope and compatibility limits.
"""
from __future__ import annotations

import base64
from io import BytesIO
import statistics

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from packages.document_ir.models import ContentWarning, DocumentIR, TextRun

FONT_ALIASES = {'Helvetica': 'Arial', 'Times-Roman': 'Times New Roman',
                'Times': 'Times New Roman', 'Courier': 'Courier New'}


def _element(name, **attrs):
    result = OxmlElement(name)
    for key, value in attrs.items():
        result.set(qn('w:' + key), str(value))
    return result


def _twips(value):
    return str(round(value * 20))


def _font_name(value):
    if not value:
        return 'Arial'
    base = value.split('-')[0]
    return FONT_ALIASES.get(value, FONT_ALIASES.get(base, value))


def _runs(paragraph, runs, content):
    for source in runs or [TextRun(content=content)]:
        run = paragraph.add_run(source.content)
        run.font.name = _font_name(source.font_name)
        run.font.size = Pt(max(1, min(200, source.font_size or 10)))
        run.bold = source.bold
        run.italic = source.italic
        if source.color and len(source.color) == 6:
            run.font.color.rgb = RGBColor.from_string(source.color)


def _paragraph_format(paragraph, size=10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(size * 1.15)
    fmt.widow_control = False
    fmt.keep_with_next = False
    fmt.keep_together = False
    paragraph._p.get_or_add_pPr().append(_element('w:snapToGrid', val=0))


def _positioned_paragraph(output, line, page):
    paragraph = output.add_paragraph()
    size = max((r.font_size or 10 for r in line.runs if r.content.strip()), default=10)
    _paragraph_format(paragraph, size)
    box = line.location
    # PDF glyph bounds omit the line's leading; Word positions its line box.
    frame = _element('w:framePr', w=_twips(min(page.width - box.x0, box.width + 5)),
        x=_twips(max(0, box.x0)), y=_twips(max(0, box.y0 - size * .18)),
        hAnchor='page', vAnchor='page', wrap='none', hRule='auto',
        hSpace=0, vSpace=0)
    paragraph._p.get_or_add_pPr().append(frame)
    _runs(paragraph, line.runs, line.content)


def _picture(anchor, image):
    if not image.data_uri or not image.data_uri.startswith('data:image/png;base64,'):
        raise ValueError('Word images must be self-contained sanitized PNG assets')
    data = base64.b64decode(image.data_uri.split(',', 1)[1], validate=True)
    if len(data) > 12 * 1024 * 1024:
        raise ValueError('Word image exceeds size limit')
    box = image.location
    if box.width <= 0 or box.height <= 0:
        raise ValueError('Word image has invalid dimensions')
    picture = anchor.add_run().add_picture(BytesIO(data), width=Pt(box.width), height=Pt(box.height))
    inline = picture._inline
    drawing = inline.getparent()
    floating = OxmlElement('wp:anchor')
    for key, val in dict(distT=0, distB=0, distL=0, distR=0, simplePos=0,
                         relativeHeight=0 if image.kind == 'graphics' else 1,
                         behindDoc=1, locked=0, layoutInCell=1, allowOverlap=1).items():
        floating.set(key, str(val))
    simple = OxmlElement('wp:simplePos')
    simple.set('x', '0')
    simple.set('y', '0')
    floating.append(simple)
    for axis, offset in [('H', box.x0), ('V', box.y0)]:
        position = OxmlElement('wp:position' + axis)
        position.set('relativeFrom', 'page')
        value = OxmlElement('wp:posOffset')
        value.text = str(round(offset * 12700))  # PDF point -> EMU
        position.append(value)
        floating.append(position)
    for child in list(inline):
        if child.tag in (qn('wp:extent'), qn('wp:effectExtent')):
            floating.append(child)
    floating.append(OxmlElement('wp:wrapNone'))
    for child in list(inline):
        floating.append(child)
    floating.find(qn('wp:docPr')).set('descr', f'{image.kind}; source page {image.page_number}')
    drawing.replace(inline, floating)


def _complex_table(table):
    # PDF panels (e.g. a receipt's two large side-by-side boxes) are often
    # detected as 1-row tables. Keep their measured text positions rather than
    # reflowing a whole panel into one long Word cell.
    return bool(table.text_lines and any(len(c.content.splitlines()) > 3 for c in table.cells))


def _table(output, source, positioned, graphics_present):
    table = output.add_table(rows=source.rows, cols=source.columns)
    table.autofit = False
    widths = source.column_widths or [source.location.width / source.columns] * source.columns
    if len(widths) != source.columns:
        widths = [source.location.width / source.columns] * source.columns
    for column, width in zip(table.columns, widths):
        column.width = Pt(width)
    properties = table._tbl.tblPr
    # Word otherwise inherits a 5.4 pt global cell margin. Its floating
    # table position and per-row padding can then differ from PDF geometry.
    default_margins = _element('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        default_margins.append(_element('w:' + side, w=0, type='dxa'))
    properties.append(default_margins)
    table_width = properties.find(qn('w:tblW'))
    table_width.set(qn('w:w'), _twips(sum(widths)))
    table_width.set(qn('w:type'), 'dxa')
    if positioned:
        properties.append(_element('w:tblpPr', horzAnchor='page', vertAnchor='page',
            tblpX=_twips(source.location.x0), tblpY=_twips(source.location.y0),
            leftFromText=0, rightFromText=0, topFromText=0, bottomFromText=0))
        properties.append(_element('w:tblOverlap', val='overlap'))
    if not graphics_present:
        table.style = 'Table Grid'
    else:
        borders = _element('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(_element('w:' + side, val='nil'))
        properties.append(borders)
    # Merge before filling, so merged placeholders never duplicate content.
    for cell in source.cells:
        if cell.row_span > 1 or cell.col_span > 1:
            table.cell(cell.row_index, cell.col_index).merge(table.cell(
                min(source.rows - 1, cell.row_index + cell.row_span - 1),
                min(source.columns - 1, cell.col_index + cell.col_span - 1)))
    occupied = set()
    for source_cell in source.cells:
        coord = (source_cell.row_index, source_cell.col_index)
        if coord in occupied:
            continue
        for r in range(source_cell.row_index, min(source.rows, source_cell.row_index + source_cell.row_span)):
            for c in range(source_cell.col_index, min(source.columns, source_cell.col_index + source_cell.col_span)):
                occupied.add((r, c))
        cell = table.cell(*coord)
        cell.width = Pt(sum(widths[source_cell.col_index:source_cell.col_index + source_cell.col_span]))
        cell.text = ''
        size = max((r.font_size or 10 for r in source_cell.runs if r.content.strip()), default=10) if source_cell.content else 1
        margins = _element('w:tcMar')
        for side in ('top', 'left', 'bottom', 'right'):
            margins.append(_element('w:' + side, w=0, type='dxa'))
        cell._tc.get_or_add_tcPr().append(margins)
        paragraph = cell.paragraphs[0]
        _paragraph_format(paragraph, size)
        paragraph.paragraph_format.left_indent = Pt(source_cell.padding_left if source_cell.padding_left is not None else 2)
        # Use paragraph spacing for each cell. Word takes the maximum top
        # cell margin across a row, incorrectly shifting two-line cells when
        # a one-line neighbour is vertically centred in the source PDF.
        paragraph.paragraph_format.space_before = Pt(max(0, (source_cell.padding_top or 0) - size * .18))
        if source_cell.original_source and source.text_lines:
            box = source_cell.original_source
            tops = sorted({line.location.y0 for line in source.text_lines
                           if box.x0 <= (line.location.x0 + line.location.x1) / 2 <= box.x1 and
                           box.y0 <= (line.location.y0 + line.location.y1) / 2 <= box.y1})
            distances = [b - a for a, b in zip(tops, tops[1:]) if size * .5 <= b - a <= size * 3]
            if distances:
                paragraph.paragraph_format.line_spacing = Pt(statistics.median(distances))
        if not source_cell.content:
            # PDF grids often contain padding rows only 3-5 pt high. A default
            # 10 pt empty Word paragraph would enlarge every padding row and
            # push later text over the next floating table or heading.
            paragraph.paragraph_format.line_spacing = Pt(1)
            mark = _element('w:rPr')
            mark.append(_element('w:sz', val=2))
            paragraph._p.get_or_add_pPr().append(mark)
        _runs(paragraph, source_cell.runs, source_cell.content)
    for row, height in zip(table.rows, source.row_heights):
        row.height = Pt(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST  # Never clip edited text.
        row._tr.get_or_add_trPr().append(_element('w:cantSplit'))


def _add_warning(page, code, message):
    if not any(w.code == code for w in page.warnings):
        page.warnings.append(ContentWarning(level='warning', code=code, message=message,
                                            page_number=page.page_number))


def export_docx(document: DocumentIR, *, include_warnings: bool = False) -> bytes:
    output = Document()
    output.core_properties.comments = "F0 prototype; source SHA256: " + document.source_hash
    normal = output.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    for index, page in enumerate(document.pages):
        section = output.add_section(WD_SECTION_START.NEW_PAGE) if index else output.sections[0]
        section.page_width = Pt(page.width)
        section.page_height = Pt(page.height)
        section.orientation = WD_ORIENT.LANDSCAPE if page.width > page.height else WD_ORIENT.PORTRAIT
        section.header_distance = section.footer_distance = Pt(0)
        section.top_margin = section.bottom_margin = Pt(1)
        section.left_margin = section.right_margin = Pt(1)
        positioned = page.metadata.get('word_layout') == 'positioned-v1'
        if not positioned:
            section.top_margin = section.bottom_margin = Pt(24)
            section.left_margin = section.right_margin = Pt(24)
        anchor = output.add_paragraph()
        _paragraph_format(anchor, 1)
        for image in sorted(page.images, key=lambda image: image.kind != 'graphics'):
            _picture(anchor, image)
        items = list(page.blocks)
        for table in page.tables:
            if positioned and _complex_table(table):
                items.extend(table.text_lines)
                _add_warning(page, 'DOCX_TABLE_AS_POSITIONED_TEXT',
                    'A complex PDF panel is represented as positioned editable text, not as a Word table.')
            else:
                items.append(table)
        for item in sorted(items, key=lambda b: (b.location.y0, b.location.x0)):
            if hasattr(item, 'cells'):
                _table(output, item, positioned, any(i.kind == 'graphics' for i in page.images))
            elif positioned:
                _positioned_paragraph(output, item, page)
            else:
                paragraph = output.add_paragraph()
                _paragraph_format(paragraph)
                _runs(paragraph, item.runs, item.content)
        runs = [r for b in page.blocks for r in b.runs]
        runs += [r for t in page.tables for c in t.cells for r in c.runs]
        substitutions = sorted({r.font_name for r in runs if r.font_name and _font_name(r.font_name) != r.font_name})
        if substitutions:
            _add_warning(page, 'DOCX_FONT_SUBSTITUTED',
                         'Standard PDF fonts mapped to Word fonts: ' + ', '.join(substitutions))
    if include_warnings:
        output.add_page_break()
        output.add_heading('Conversion warnings', level=1)
        for issue in document.warnings + [w for p in document.pages for w in p.warnings]:
            output.add_paragraph(f'Page {issue.page_number}: {issue.code} — {issue.message}')
    buffer = BytesIO()
    output.save(buffer)
    return buffer.getvalue()
